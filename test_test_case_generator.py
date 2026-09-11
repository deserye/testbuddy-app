from io import BytesIO

from openpyxl import load_workbook

from test_case_generator import cases_to_excel, cases_to_json, cases_to_markdown, cases_to_rows, direct_input_candidates, generate_test_cases, reconciliation_scenarios

text = "The reconciliation service shall match ESHTRN and ACTRN using the agreed identifier and amount rules. It must flag duplicate entries and mismatched amounts for review."
candidates = direct_input_candidates(text, input_type="Business requirement", scenario="Mismatched amount")
assert candidates
assert candidates[0]["source_document"] == "Direct Business requirement input"
assert "Mismatched amount" in candidates[0]["source_section"]

cases, mode = generate_test_cases(candidates, client=None, include_negative=True, max_cases=4)
assert cases
assert len(cases) <= 4
assert any("negative" in case.test_type.lower() or "validation" in case.test_type.lower() for case in cases)
assert all(case.source_document for case in cases)
assert all(case.steps for case in cases)
assert "TC-" in cases_to_markdown(cases)
assert "test_case_id" in cases_to_json(cases)
assert cases_to_rows(cases)[0]["steps"]
assert "Duplicate entry" in reconciliation_scenarios()
workbook = load_workbook(BytesIO(cases_to_excel(cases)))
assert workbook.sheetnames == ["Test Cases", "Steps"]
assert workbook["Test Cases"]["A2"].value == cases[0].test_case_id
assert workbook["Steps"]["A2"].value == cases[0].test_case_id
print("TestBuddy test-case generator smoke tests passed.")
